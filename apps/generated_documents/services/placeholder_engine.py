"""
DOCX placeholder replacement engine (Step 4.2 of document generation).

Given a ``.docx`` template and the validated data context produced by Step 4.1
(``apps.generated_documents.services.context``), this service writes a new
``.docx`` document in which every approved ``{{ ... }}`` placeholder is
replaced by its actual value. The source template is never modified: it is
read into memory, the in-memory copy is edited, and a separate output document
is produced.

Security
--------
Only the *controlled* variable catalog from Step 3
(``apps.document_templates.variables``) is ever resolved. Placeholder keys are
normalized text (``namespace.field``) and looked up in the context dict; there
is no ``eval``, no attribute traversal, and no template-expression execution.
Unknown/malformed placeholders and missing *required* values raise a clear
error instead of producing a silently wrong document.

Placeholder syntax
------------------
    {{ company.name }}
    {{ client.voen }}
    {{ document.number }}

Whitespace and casing around the key are normalized (``{{client.name}}`` and
``{{   Company.Name   }}`` resolve to the same variable). Only the
``{{ ... }}`` form is recognized.

Repeating line-item sections (Step 4.3)
---------------------------------------
A template may repeat a block once per document line item using the controlled
``{{#each items}}`` ... ``{{/each}}`` syntax. ``items`` is the only supported
collection name and maps to the document's line items.

* **Table rows (priority use case):** the prototype row's *first* cell starts
  with ``{{#each items}}`` and its *last* cell ends with ``{{/each}}``, and the
  cells contain ``{{ item.* }}`` placeholders. The row is cloned once per item
  (zero items removes the row entirely) and the markers are stripped. The
  header/totals rows are left untouched.
* **Paragraphs:** ``{{#each items}}`` and ``{{/each}}`` on their own
  paragraphs enclose a block of paragraphs that is replicated per item (a
  single row or cell may contain such a block).

Supported namespaces (authoritative catalog in ``variables.py``)
----------------------------------------------------------------
    company.*    company.name, company.legal_name, company.voen,
                 company.registration_number
    client.*     client.name, client.voen, client.email, client.phone
    document.*   document.number, document.date, document.type
    item.*       (inside {{#each items}} only) item.description,
                 item.quantity, item.unit_price, item.total, ... — resolved
                 from the current line item.

Nested repeating sections, conditionals, filters, function calls and
expressions are deliberately NOT supported. ``items`` is a fixed collection;
``{{#each unknown_collection}}`` is rejected.

DOCX locations processed
------------------------
* normal body paragraphs
* table cells / rows (including nested tables)
* existing headers/footers (only parts that already exist — no empty
  header/footer parts are created)

Split runs
----------
DOCX text is frequently split across multiple runs inside one paragraph
(e.g. ``{{ client`` + ``.name }}``). The engine concatenates the paragraph's
run text to locate placeholders and rewrites only the affected runs, so a
placeholder split across runs is still replaced while the surrounding
formatting is preserved.

Formatting preservation
-----------------------
When a table row is cloned, the clone is a deep copy of the prototype row's
XML, so cell/paragraph/font formatting, borders and alignment are preserved.
Only the marker text is stripped and the ``{{ ... }}`` placeholders replaced.

Value formatting
----------------
Values are converted to text deterministically: strings pass through,
``None`` becomes the empty string, ``date``/``datetime`` become an ISO date
(``YYYY-MM-DD``), and ``Decimal`` is rendered without trailing zeros. No
localization/currency formatting is applied (deferred to a later step).

Errors
------
* :class:`UnknownPlaceholderError` — a placeholder is malformed or does not
  match any approved variable (e.g. ``{{ client.unknown_field }}``,
  ``{{ some_random_variable }}``, ``{{#each bogus}}``), a repeating section
  is unbalanced or nested, or an ``item.*`` variable is used outside a
  repeating section. The engine never replaces or executes it.
* :class:`MissingRequiredValueError` — an approved *required* variable
  (``required=True`` in the catalog) has no value in the context.
* Optional variables whose value is missing/``None`` resolve to the empty
  string and never crash generation.

The engine validates the whole document before replacing anything, so all
problems are reported at once and no output is produced when validation fails.
"""

import os
import re
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from apps.document_templates.variables import (
    EACH_COLLECTION,
    PLACEHOLDER_RE,
    VARIABLE_BY_KEY,
    _normalize_placeholder,
)

EACH_RE = re.compile(r"\{\{\s*#each\s+([a-z][a-z0-9_]*)\s*\}\}")
END_EACH_RE = re.compile(r"\{\{\s*/each\s*\}\}")

ITEM_VARIABLE_KEYS = frozenset(k for k in VARIABLE_BY_KEY if k.startswith("item."))

_MISSING = object()


class PlaceholderEngineError(Exception):
    """Base error for DOCX placeholder replacement failures."""


class UnknownPlaceholderError(PlaceholderEngineError):
    """A placeholder is malformed or does not match any approved variable."""


class MissingRequiredValueError(PlaceholderEngineError):
    """An approved required variable has no resolvable value."""


# ---------------------------------------------------------------------------
# Value formatting
# ---------------------------------------------------------------------------


def _format_decimal(value):
    """Render a Decimal without trailing zeros and without exponent notation."""
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def format_value(value):
    """
    Convert a single context value to its deterministic text representation.

    ``None`` becomes the empty string (the project's empty-value convention).
    Dates/datetimes render as an ISO date (``YYYY-MM-DD``), not the raw Python
    repr, so ``2026-09-01`` does not unexpectedly become
    ``2026-09-01 00:00:00``.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal):
        return _format_decimal(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _resolve(key, context, item):
    """
    Resolve an approved ``namespace.field`` key against ``context``.

    ``item.*`` keys resolve against the current repeating-section item; other
    keys resolve against the top-level base context. ``item`` is ``None``
    outside a repeating section.
    """
    if key in ITEM_VARIABLE_KEYS:
        if item is None:
            raise UnknownPlaceholderError(
                f"'{{{{ {key} }}}}' may only be used inside a repeating section."
            )
        field = key.split(".", 1)[1]
        value = item.get(field, _MISSING)
        if value is _MISSING or value is None:
            return ""
        return format_value(value)

    variable = VARIABLE_BY_KEY[key]
    namespace, _, field = key.partition(".")
    value = context.get(namespace, {}).get(field, _MISSING)
    if value is _MISSING or value is None:
        if variable.required:
            raise MissingRequiredValueError(
                f"Required template variable '{{{{ {key} }}}}' has no value "
                f"in the generation context."
            )
        return ""
    return format_value(value)


def _paragraph_text(p_el):
    """Concatenated text of a ``<w:p>`` element (all its runs)."""
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _apply_spans(runs, full_text, spans):
    """
    Rewrite ``runs`` given ``spans`` = list of (start, end, replacement)
    offsets into the concatenated run text. Preserves all text outside the
    spans and each run's formatting.
    """
    offsets = []
    cursor = 0
    for run in runs:
        offsets.append(cursor)
        cursor += len(run.text or "")

    for index, run in enumerate(runs):
        text = run.text or ""
        if not text:
            continue
        run_start = offsets[index]
        run_end = run_start + len(text)

        pieces = []
        pos = run_start
        for span_start, span_end, replacement in spans:
            if span_end <= run_start or span_start >= run_end:
                continue
            keep_end = min(span_start, run_end)
            if keep_end > pos:
                pieces.append(text[pos - run_start : keep_end - run_start])
            if span_start >= run_start:
                pieces.append(replacement)
            pos = max(pos, min(span_end, run_end))
        if pos < run_end:
            pieces.append(text[pos - run_start : run_end - run_start])
        run.text = "".join(pieces)


def _replace_paragraph(paragraph, context, item):
    """
    Replace every placeholder in one paragraph.

    The paragraph's run texts are concatenated so a placeholder split across
    runs is found; each affected run is then rewritten in place, preserving
    the formatting of the surrounding text.
    """
    runs = paragraph.runs
    if not runs:
        return
    full_text = "".join(r.text or "" for r in runs)
    if "{{" not in full_text:
        return
    matches = list(PLACEHOLDER_RE.finditer(full_text))
    if not matches:
        return
    spans = [
        (m.start(), m.end(), _resolve(_normalize_placeholder(m.group(1)), context, item))
        for m in matches
    ]
    _apply_spans(runs, full_text, spans)


def _strip_token(p_el, regex):
    """Remove every ``regex`` match (e.g. an each/endeach marker) from a paragraph."""
    para = Paragraph(p_el, None)
    runs = para.runs
    if not runs:
        return
    full_text = "".join(r.text or "" for r in runs)
    if not full_text:
        return
    matches = list(regex.finditer(full_text))
    if not matches:
        return
    spans = [(m.start(), m.end(), "") for m in matches]
    _apply_spans(runs, full_text, spans)


def _cell_first_text(tr):
    tcs = tr.findall(qn("w:tc"))
    if not tcs:
        return ""
    return _paragraph_text(tcs[0])


def _cell_last_text(tr):
    tcs = tr.findall(qn("w:tc"))
    if not tcs:
        return ""
    return _paragraph_text(tcs[-1])


def _is_prototype_row(tr):
    return bool(
        EACH_RE.search(_cell_first_text(tr))
        and END_EACH_RE.search(_cell_last_text(tr))
    )


def _strip_row_markers(tr):
    """Remove the ``{{#each items}}`` / ``{{/each}}`` markers from first/last cells."""
    tcs = tr.findall(qn("w:tc"))
    if not tcs:
        return
    for p_el in tcs[0].findall(qn("w:p")):
        _strip_token(p_el, EACH_RE)
    last = tcs[-1]
    if last is not tcs[0]:
        for p_el in last.findall(qn("w:p")):
            _strip_token(p_el, END_EACH_RE)
        return
    for p_el in tcs[0].findall(qn("w:p")):
        _strip_token(p_el, END_EACH_RE)


def _expand_paragraph_sections(p_els, context, items, ambient_item):
    """
    Given a list of ``<w:p>`` elements, expand any ``{{#each items}}`` ...
    ``{{/each}}`` paragraph sections and replace placeholders.

    Returns ``(new_els, counts)`` where ``new_els`` is the flattened list of
    rendered ``<w:p>`` elements and ``counts[i]`` says how many of them were
    produced for original paragraph ``p_els[i]`` (0 for dropped marker/empty
    paragraphs, >1 for replicated section bodies). The two are aligned so
    callers can rebuild containers while preserving non-paragraph elements.
    """
    out = []
    counts = []
    i = 0
    n = len(p_els)
    while i < n:
        el = p_els[i]
        m = EACH_RE.search(_paragraph_text(el))
        if m:
            col = m.group(1)
            if col != EACH_COLLECTION:
                _replace_paragraph(Paragraph(el, None), context, ambient_item)
                counts.append(1)
                out.append(el)
                i += 1
                continue
            end = None
            j = i + 1
            while j < n:
                if END_EACH_RE.search(_paragraph_text(p_els[j])):
                    end = j
                    break
                j += 1
            if end is None:
                _replace_paragraph(Paragraph(el, None), context, ambient_item)
                counts.append(1)
                out.append(el)
                i += 1
                continue
            section = p_els[i + 1 : end]
            if not items:
                for _ in p_els[i : end + 1]:
                    counts.append(0)
                i = end + 1
                continue
            for item in items:
                for sp in section:
                    clone = deepcopy(sp)
                    _replace_paragraph(Paragraph(clone, None), context, item)
                    out.append(clone)
            counts.append(0)  # opener marker dropped
            for _ in section:
                counts.append(len(items))
            counts.append(0)  # closer marker dropped
            i = end + 1
            continue
        _replace_paragraph(Paragraph(el, None), context, ambient_item)
        counts.append(1)
        out.append(el)
        i += 1
    return out, counts


def _replicate_paragraph_sections(container, context, items, ambient_item):
    """
    Replace the ``<w:p>`` children of ``container`` (a ``w:body``, ``w:tc``,
    ``w:hdr`` or ``w:ftr``), expanding repeating paragraph sections and
    replacing placeholders. Non-paragraph children (tables, section props)
    keep their relative positions.
    """
    p_els = [el for el in container if el.tag == qn("w:p")]
    new_els, counts = _expand_paragraph_sections(p_els, context, items, ambient_item)
    new_children = []
    new_index = 0
    for child in list(container):
        if child.tag == qn("w:p"):
            index = p_els.index(child)
            count = counts[index]
            new_children.extend(new_els[new_index : new_index + count])
            new_index += count
        else:
            new_children.append(child)
    for child in list(container):
        container.remove(child)
    for child in new_children:
        container.append(child)


def _fill_row(tr, context, item, items):
    """Replace placeholders and expand sections inside one table row's cells."""
    for tc in tr.findall(qn("w:tc")):
        _replicate_paragraph_sections(tc, context, items, item)
        for tbl_el in tc.findall(qn("w:tbl")):
            _process_table(Table(tbl_el, None), context, items, item)


def _process_table(table, context, items, ambient_item=None):
    """
    Process one table: replicate any ``{{#each items}}`` prototype rows, and
    replace placeholders / expand sections in every row's cells. Nested tables
    are processed recursively with the ambient item scope.
    """
    tbl = table._tbl
    tr_els = tbl.findall(qn("w:tr"))
    has_marker = any(_is_prototype_row(tr) for tr in tr_els)
    if has_marker:
        _replicate_rows(tbl, tr_els, context, items)
        return
    for tr in tr_els:
        _fill_row(tr, context, ambient_item, items)


def _replicate_rows(tbl, tr_els, context, items):
    """Replicate prototype rows once per item (drop them when items is empty)."""
    replacement = []
    for tr in tr_els:
        if _is_prototype_row(tr):
            if not items:
                continue
            for item in items:
                clone = deepcopy(tr)
                _strip_row_markers(clone)
                _fill_row(clone, context, item, items)
                replacement.append(clone)
            continue
        _fill_row(tr, context, None, items)
        replacement.append(tr)
    for tr in tr_els:
        tbl.remove(tr)
    for tr in replacement:
        tbl.append(tr)


def _existing_header_footer_parts(document):
    """
    Yield only the header/footer parts a document already references.

    The python-docx high-level ``section.header`` / ``section.footer`` API
    *creates* an empty part when accessed; iterating the XML references
    instead means a template with no headers/footers is not given any.
    """
    part = document.part
    seen = set()
    for sect_pr in part.element.body.iter(qn("w:sectPr")):
        for tag in ("w:headerReference", "w:footerReference"):
            for ref in sect_pr.findall(qn(tag)):
                related_id = ref.get(qn("r:id"))
                if not related_id:
                    continue
                related = part.related_parts.get(related_id)
                if related is None or id(related) in seen:
                    continue
                seen.add(id(related))
                yield related


def _validate_paragraph_container(container, context, in_section, acc, ignore_markers=False):
    """
    Validate markers and placeholders within one paragraph container
    (``w:body``, ``w:tc``, or header/footer root).

    ``ignore_markers`` is set for cells that hold a *row-level* marker, so the
    ``{{#each items}}`` / ``{{/each}}`` tokens there are not double-counted as
    paragraph-level markers.
    """
    unknown, missing = acc
    p_els = [el for el in container if el.tag == qn("w:p")]
    depth = 1 if in_section else 0
    for el in p_els:
        text = _paragraph_text(el)
        if not ignore_markers:
            for m in EACH_RE.finditer(text):
                col = m.group(1)
                if col != EACH_COLLECTION:
                    unknown.append(
                        f"Unsupported repeating collection '{{{{#each {col}}}}}'. "
                        f"Only 'items' is supported."
                    )
                depth += 1
                if depth > 1:
                    unknown.append("Nested repeating sections are not supported.")
            if END_EACH_RE.search(text):
                depth -= 1
                if depth < 0:
                    unknown.append(
                        "Found '{{/each}}' without a matching '{{#each ...}}'."
                    )
        for pm in PLACEHOLDER_RE.finditer(text):
            raw = pm.group(1)
            stripped = raw.strip()
            if stripped.startswith("#each") or stripped.startswith("/each"):
                continue
            key = _normalize_placeholder(raw)
            if key is None or key not in VARIABLE_BY_KEY:
                unknown.append(
                    f"Unsupported placeholder '{{{{{raw}}}}}' does not match "
                    f"any approved template variable."
                )
                continue
            if key in ITEM_VARIABLE_KEYS and depth == 0:
                unknown.append(
                    f"'{{{{ {key} }}}}' may only be used inside a repeating section."
                )
                continue
            variable = VARIABLE_BY_KEY[key]
            namespace, _, field = key.partition(".")
            value = context.get(namespace, {}).get(field, _MISSING)
            if variable.required and (value is _MISSING or value is None):
                missing.append(
                    f"Required template variable '{{{{ {key} }}}}' has no "
                    f"value in the generation context."
                )
    if not ignore_markers and not in_section and depth != 0:
        unknown.append(
            "A '{{#each ...}}' is missing its matching '{{/each}}'."
        )


def _validate_table_el(tbl_el, context, acc):
    """Validate repeating row sections and cell contents within one table."""
    unknown, missing = acc
    tr_els = tbl_el.findall(qn("w:tr"))
    depth = 0
    for tr in tr_els:
        first = _cell_first_text(tr)
        last = _cell_last_text(tr)
        em = EACH_RE.search(first)
        cm = END_EACH_RE.search(last)
        if em:
            col = em.group(1)
            if col != EACH_COLLECTION:
                unknown.append(
                    f"Unsupported repeating collection '{{{{#each {col}}}}}'. "
                    f"Only 'items' is supported."
                )
            depth += 1
            if depth > 1:
                unknown.append("Nested repeating sections are not supported.")
        if cm:
            depth -= 1
            if depth < 0:
                unknown.append(
                    "Found '{{/each}}' without a matching '{{#each ...}}'."
                )
        # A prototype row both opens and closes its own (single-row) section,
        # so its cells are within a section even though the running depth has
        # already returned to zero.
        in_section = depth > 0 or em is not None
        tcs = tr.findall(qn("w:tc"))
        for index, tc in enumerate(tcs):
            ignore = (
                index == 0 and em is not None
                or index == len(tcs) - 1 and len(tcs) > 1 and cm is not None
            )
            _validate_paragraph_container(
                tc, context, in_section, acc, ignore_markers=ignore
            )
            for nested in tc.findall(qn("w:tbl")):
                _validate_table_el(nested, context, acc)
    if depth != 0:
        unknown.append("A repeating row section is missing its matching '{{/each}}'.")


def _validate(document, context):
    """
    Validate every placeholder and repeating section in the document before
    any replacement.

    Raises :class:`UnknownPlaceholderError` (malformed/unknown keys, unknown
    or unbalanced/nested repeating sections, ``item.*`` outside a section) or
    :class:`MissingRequiredValueError` (approved required variable without a
    value). All problems across the whole document are reported at once.
    """
    unknown = []
    missing = []
    body = document.element.body
    _validate_paragraph_container(body, context, False, (unknown, missing))
    for tbl_el in [c for c in body if c.tag == qn("w:tbl")]:
        _validate_table_el(tbl_el, context, (unknown, missing))
    for part in _existing_header_footer_parts(document):
        _validate_paragraph_container(part.element, context, False, (unknown, missing))

    if unknown:
        raise UnknownPlaceholderError("\n".join(unknown))
    if missing:
        raise MissingRequiredValueError("\n".join(missing))


def _replace(document, context):
    """Replace placeholders and expand repeating sections throughout the document."""
    items = context.get("items", [])
    body = document.element.body
    _replicate_paragraph_sections(body, context, items, None)
    for tbl_el in [c for c in body if c.tag == qn("w:tbl")]:
        _process_table(Table(tbl_el, None), context, items, None)
    for part in _existing_header_footer_parts(document):
        _replicate_paragraph_sections(part.element, context, items, None)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def _load_document(template_file):
    """Open a .docx from a path or binary file-like object."""
    try:
        if isinstance(template_file, (str, os.PathLike)):
            return Document(str(template_file))
        if hasattr(template_file, "read"):
            content = template_file.read()
            try:
                template_file.seek(0)
            except (AttributeError, OSError):
                pass
            return Document(BytesIO(content))
    except Exception as exc:
        raise PlaceholderEngineError(
            "The template is not a readable .docx file."
        ) from exc
    raise PlaceholderEngineError(
        "template_file must be a path or a binary file-like object."
    )


def _save_document(document, output_file):
    """Persist the in-memory document to a path or binary file-like object."""
    if isinstance(output_file, (str, os.PathLike)):
        document.save(str(output_file))
        return output_file
    if hasattr(output_file, "write"):
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        output_file.write(buffer.read())
        return output_file
    raise PlaceholderEngineError(
        "output_file must be a path or a binary file-like object."
    )


def render_document(template_file, output_file, context):
    """
    Render a .docx template with ``context``, producing a new document.

    Args:
        template_file: path or binary file-like object for the source .docx.
            It is read-only and never modified.
        output_file: path or binary file-like object to write the result to.
        context: the controlled template-variable context from Step 4.1, e.g.
            ``{"company": {...}, "client": {...}, "document": {...}}``.

    Returns:
        ``output_file``.

    Raises:
        UnknownPlaceholderError: the template contains a malformed or unknown
            placeholder (nothing is written).
        MissingRequiredValueError: an approved required variable has no value
            (nothing is written).
        PlaceholderEngineError: the input is not a readable .docx.
    """
    document = _load_document(template_file)
    _validate(document, context)
    _replace(document, context)
    return _save_document(document, output_file)
