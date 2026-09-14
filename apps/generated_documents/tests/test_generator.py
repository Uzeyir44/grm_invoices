import os
import shutil
import tempfile
import zipfile
from decimal import Decimal
from io import BytesIO

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from docx import Document
from lxml import etree

from apps.authentication.tests.test_authentication import make_company
from apps.clients.models import (
    Client,
    ClientLegalEntity,
)
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import (
    DocumentLineItem,
    GeneratedDocument,
    GeneratedDocumentVersion,
)
from apps.generated_documents.services.context import GenerationContextError
from apps.generated_documents.services.generator import (
    GenerationError,
    generate_document,
    select_template_version,
)
from apps.generated_documents.services.placeholder_engine import (
    UnknownPlaceholderError,
)
from apps.reference_data.models import Currency, DocumentType

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_generation_tests_")


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class GenerationServiceTestCase(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.company = make_company()
        self.document_type, _ = DocumentType.objects.get_or_create(
            code="invoice", defaults={"name": "Invoice"}
        )
        self.template = DocumentTemplate.objects.create(
            company=self.company,
            document_type=self.document_type,
            name="Invoice Template",
        )
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        self.client = self._make_legal_client(self.company)
        self.template_version = self._make_template_version()
        self.document = self._make_document(clients=self.client)

    # ------------------------------------------------------------------ setup

    def _make_legal_client(self, company):
        client = Client.objects.create(
            company=company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client,
            company=company,
            legal_name="Acme Ltd",
            voen="1234567890",
        )
        return client

    def _make_template_version(self, file_content=None, active=True, version_number=1):
        if file_content is None:
            file_content = build_invoice_template()
        tv = TemplateVersion.objects.create(
            template=self.template,
            company=self.company,
            version_number=version_number,
            storage_path=f"v{version_number}",
            is_active=active,
        )
        tv.file.save("tpl.docx", ContentFile(file_content), save=True)
        return tv

    def _make_document(self, clients, number="INV-001", template=None):
        return GeneratedDocument.objects.create(
            company=self.company,
            client=clients,
            document_type=self.document_type,
            template=template or self.template,
            document_number=number,
        )

    def _make_source_version(self, document, descriptions):
        source = GeneratedDocumentVersion.objects.create(
            generated_document=document,
            company=document.company,
            version_number=1,
            template_version=self.template_version,
            currency=self.currency,
            generation_snapshot={},
        )
        for index, description in enumerate(descriptions, start=1):
            DocumentLineItem.objects.create(
                generated_document_version=source,
                company=document.company,
                currency=self.currency,
                line_number=index,
                description=description,
                quantity=Decimal("2.0000"),
                unit="h",
                unit_price=Decimal("100.0000"),
                discount_percent=Decimal("0.00"),
                vat_rate_percent=Decimal("18.00"),
                subtotal=Decimal("200.0000"),
                vat_amount=Decimal("36.0000"),
                total=Decimal("236.0000"),
            )
        return source

    # ------------------------------------------------------------------ helpers

    def _docx_text(self, storage_path):
        with default_storage_file(storage_path) as fobj:
            with zipfile.ZipFile(fobj) as zf:
                xml = zf.read("word/document.xml")
        root = etree.fromstring(xml)
        return "".join(t.text or "" for t in root.iter(W_NS + "t"))

    def _read_docx_bytes(self, storage_path):
        with default_storage_file(storage_path) as fobj:
            return fobj.read()


def default_storage_file(storage_path):
    from django.core.files.storage import default_storage

    return default_storage.open(storage_path, "rb")


def build_invoice_template():
    """Small real DOCX fixture: simple placeholders + an items table."""
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    document.add_paragraph("VOEN: {{ client.voen }}")
    table = document.add_table(rows=0, cols=3)
    header = table.add_row().cells
    header[0].text = "Description"
    header[1].text = "Qty"
    header[2].text = "Total"
    proto = table.add_row().cells
    proto[0].text = "{{#each items}}"
    proto[1].text = "{{ item.description }}"
    proto[2].text = "{{/each}}"
    totals = table.add_row().cells
    totals[0].text = ""
    totals[1].text = "End"
    totals[2].text = ""
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


class BasicGenerationTests(GenerationServiceTestCase):
    def test_basic_generation_produces_docx(self):
        version = generate_document(self.document, currency=self.currency)
        self.assertTrue(version.storage_path_docx)
        self.assertTrue(os.path.exists(os.path.join(TEMP_MEDIA, version.storage_path_docx)))
        self._read_docx_bytes(version.storage_path_docx)
        # valid OOXML package with a body
        with default_storage_file(version.storage_path_docx) as fobj:
            with zipfile.ZipFile(fobj) as zf:
                self.assertIn("word/document.xml", zf.namelist())

    def test_simple_placeholders_are_replaced(self):
        version = generate_document(self.document, currency=self.currency)
        text = self._docx_text(version.storage_path_docx)
        self.assertIn("Invoice: INV-001", text)
        self.assertIn("Client: Acme Ltd", text)
        self.assertIn("VOEN: 1234567890", text)
        self.assertNotIn("{{", text)

    def test_repeating_items_integration(self):
        source = self._make_source_version(self.document, ["Consulting", "Development"])
        version = generate_document(self.document)
        self.assertIsNotNone(version)
        text = self._docx_text(version.storage_path_docx)
        self.assertIn("Consulting", text)
        self.assertIn("Development", text)
        self.assertIn("Description", text)


class VersionPersistenceTests(GenerationServiceTestCase):
    def test_generated_version_records_expected_fields(self):
        source = self._make_source_version(self.document, ["Consulting"])
        version = generate_document(
            self.document,
            exchange_rate_used=Decimal("1.70000000"),
            vat_rate_used_percent=Decimal("18.00"),
        )
        version.refresh_from_db()
        self.assertEqual(version.generated_document_id, self.document.id)
        self.assertEqual(version.company_id, self.company.id)
        self.assertEqual(version.version_number, 2)
        self.assertEqual(version.template_version_id, self.template_version.id)
        self.assertEqual(version.currency_id, self.currency.id)
        self.assertEqual(version.exchange_rate_used, Decimal("1.70000000"))
        self.assertEqual(version.vat_rate_used_percent, Decimal("18.00"))
        self.assertTrue(version.storage_path_docx.endswith(".docx"))
        self.assertEqual(version.generation_metadata["engine_version"], "4.4")
        self.assertIn("items", version.generation_snapshot)

    def test_status_advances_to_generated(self):
        self.document.refresh_from_db()
        self.assertEqual(self.document.status, GeneratedDocument.Status.DRAFT)
        version = generate_document(self.document, currency=self.currency)
        self.document.refresh_from_db()
        self.assertEqual(self.document.status, GeneratedDocument.Status.GENERATED)
        self.assertIsNotNone(version)

    def test_version_number_is_incremented_per_generation(self):
        self.assertEqual(self.document.versions.count(), 0)
        first = generate_document(self.document, currency=self.currency)
        second = generate_document(self.document, currency=self.currency)
        self.assertEqual(first.version_number, 1)
        self.assertEqual(second.version_number, 2)
        self.assertEqual(self.document.versions.count(), 2)
        self.assertNotEqual(first.storage_path_docx, second.storage_path_docx)
        self.assertTrue(second.generated_document_id == first.generated_document_id)

    def test_regeneration_keeps_prior_version_retrievable(self):
        first = generate_document(self.document, currency=self.currency)
        first_path = first.storage_path_docx
        second = generate_document(self.document, currency=self.currency)
        # The first file must still be on disk; both versions remain in the DB.
        self.assertTrue(os.path.exists(os.path.join(TEMP_MEDIA, first_path)))
        self.assertEqual(
            GeneratedDocumentVersion.objects.filter(
                generated_document=self.document
            ).count(),
            2,
        )

    def test_template_tracking_records_used_template_version(self):
        tv2 = self._make_template_version(
            file_content=build_invoice_template(), version_number=2, active=False
        )
        self.template_version.is_active = False
        self.template_version.save(update_fields=["is_active"])
        tv2.is_active = True
        tv2.save(update_fields=["is_active"])
        version = generate_document(self.document, currency=self.currency)
        self.assertEqual(version.template_version_id, tv2.id)


class TemplatePreservationTests(GenerationServiceTestCase):
    def test_source_template_file_unchanged(self):
        stored_name = self.template_version.file.name
        with default_storage_file(stored_name) as fobj:
            before = fobj.read()
        version = generate_document(self.document, currency=self.currency)
        with default_storage_file(stored_name) as fobj:
            after = fobj.read()
        self.assertEqual(after, before)
        self.assertTrue(version.storage_path_docx)


class FailureBehaviorTests(GenerationServiceTestCase):
    def test_missing_required_data_fails_without_version(self):
        doc = self._make_document(clients=None, number="INV-NONE")
        with self.assertRaises(GenerationContextError):
            generate_document(doc)
        self.assertEqual(GeneratedDocumentVersion.objects.filter(generated_document=doc).count(), 0)
        doc.refresh_from_db()
        self.assertEqual(doc.status, GeneratedDocument.Status.DRAFT)

    def test_unknown_placeholder_fails_without_version(self):
        bad = BytesIO()
        d = Document()
        d.add_paragraph("{{ bogus.field }}")
        d.save(bad)
        bad.seek(0)
        tv = self._make_template_version(
            file_content=bad.getvalue(), version_number=3, active=False
        )
        self.template_version.is_active = False
        self.template_version.save(update_fields=["is_active"])
        tv.is_active = True
        tv.save(update_fields=["is_active"])

        with self.assertRaises(UnknownPlaceholderError):
            generate_document(self.document, currency=self.currency)
        self.assertEqual(
            GeneratedDocumentVersion.objects.filter(generated_document=self.document).count(),
            0,
        )

    def test_no_active_template_version_fails_without_version(self):
        self.template_version.is_active = False
        self.template_version.save(update_fields=["is_active"])
        with self.assertRaises(GenerationError) as manager:
            generate_document(self.document, currency=self.currency)
        self.assertIn("active", str(manager.exception))
        self.assertEqual(
            GeneratedDocumentVersion.objects.filter(generated_document=self.document).count(),
            0,
        )

    def test_template_without_file_fails(self):
        self.template_version.is_active = False
        self.template_version.save(update_fields=["is_active"])
        tv = TemplateVersion.objects.create(
            template=self.template,
            company=self.company,
            version_number=3,
            storage_path="v3",
            is_active=True,
        )
        with self.assertRaises(GenerationError):
            generate_document(self.document, currency=self.currency)


class TenantSafetyTests(GenerationServiceTestCase):
    """
    The DB-level composite-FK triggers already prevent a cross-company
    GeneratedDocument from existing, so the generator's own template/company
    checks are tested directly as defense-in-depth (they matter only if the
    DB invariants were bypassed), and the DB enforcement is asserted
    end-to-end.
    """

    def test_database_blocks_cross_company_template(self):
        from django.db import IntegrityError, transaction

        other_company = make_company()
        other_type = DocumentType.objects.create(code="contract", name="Contract")
        foreign_template = DocumentTemplate.objects.create(
            company=other_company, document_type=other_type, name="Foreign Tpl"
        )
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                GeneratedDocument.objects.create(
                    company=self.company,
                    client=self.client,
                    document_type=self.document_type,
                    template=foreign_template,
                    document_number="INV-DB",
                )

    def test_database_blocks_cross_company_client(self):
        from django.db import IntegrityError, transaction

        other_company = make_company()
        foreign_client = self._make_legal_client(other_company)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                GeneratedDocument.objects.create(
                    company=self.company,
                    client=foreign_client,
                    document_type=self.document_type,
                    template=self.template,
                    document_number="INV-DBC",
                )

    def test_generator_rejects_cross_company_template_version(self):
        # An unsaved document stub (valid against the real DB invariants) is
        # combined with a foreign TemplateVersion to exercise the generator's
        # own company check.
        from apps.generated_documents.services.generator import (
            _assert_template_version_matches_document,
        )

        other_company = make_company()
        other_type = DocumentType.objects.create(code="contract", name="Contract")
        foreign_template = DocumentTemplate.objects.create(
            company=other_company, document_type=other_type, name="Foreign Tpl"
        )
        foreign_tv = TemplateVersion.objects.create(
            template=foreign_template,
            company=other_company,
            version_number=1,
            storage_path="f",
            is_active=True,
        )
        stub = GeneratedDocument(
            company=self.company,
            client=self.client,
            document_type=self.document_type,
            template=self.template,
            document_number="INV-STUB",
        )
        with self.assertRaises(GenerationError):
            _assert_template_version_matches_document(foreign_tv, stub)

    def test_generator_rejects_mismatched_template(self):
        from apps.generated_documents.services.generator import (
            _assert_template_version_matches_document,
        )

        other_type = DocumentType.objects.create(code="contract", name="Contract")
        template_b = DocumentTemplate.objects.create(
            company=self.company, document_type=other_type, name="Other"
        )
        tv_b = TemplateVersion.objects.create(
            template=template_b,
            company=self.company,
            version_number=1,
            storage_path="b",
            is_active=True,
        )
        stub = GeneratedDocument(
            company=self.company,
            client=self.client,
            document_type=self.document_type,
            template=self.template,
            document_number="INV-STUB2",
        )
        with self.assertRaises(GenerationError):
            _assert_template_version_matches_document(tv_b, stub)

    def test_matching_template_version_passes_check(self):
        from apps.generated_documents.services.generator import (
            _assert_template_version_matches_document,
        )

        stub = GeneratedDocument(
            company=self.company,
            client=self.client,
            document_type=self.document_type,
            template=self.template,
            document_number="INV-OK",
        )
        # No exception means the check passed.
        _assert_template_version_matches_document(self.template_version, stub)


class FileNamingTests(GenerationServiceTestCase):
    def test_document_number_paths_are_sanitized(self):
        doc = self._make_document(clients=self.client, number="INV/../-1")
        version = generate_document(doc, currency=self.currency)
        self.assertNotIn("..", version.storage_path_docx)
        self.assertTrue(version.storage_path_docx.endswith(".docx"))


class TemplateVersionSelectionTests(GenerationServiceTestCase):
    def test_select_template_version_uses_active_version(self):
        self.assertEqual(select_template_version(self.document), self.template_version)

    def test_select_template_version_returns_none_when_no_active(self):
        self.template_version.is_active = False
        self.template_version.save(update_fields=["is_active"])
        self.assertIsNone(select_template_version(self.document))