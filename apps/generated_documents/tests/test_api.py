"""
Step 4.6 focused API tests: generation endpoint, version listing, and secure
DOCX/PDF downloads (tenant isolation + RBAC + controlled failures).

Deliberately does not rewrite the existing service tests — those still cover
the generation pipeline. These tests exercise the adapter layer only.
"""

import os
import shutil
import tempfile
import uuid
from io import BytesIO
from unittest import mock, skipUnless

from django.core.files.base import ContentFile
from django.test import override_settings
from django.urls import reverse
from docx import Document

from rest_framework import status
from rest_framework.test import APITestCase

from apps.authentication.models import Permission, Role, UserRole
from apps.authentication.tests.test_authentication import make_company, make_user
from apps.clients.models import Client, ClientLegalEntity
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import GeneratedDocument, GeneratedDocumentVersion
from apps.generated_documents.services.generator import (
    GenerationError,
    generate_document,
)
from apps.generated_documents.services.pdf_generator import (
    _find_soffice,
    PdfGenerationError,
    generate_pdf,
)
from apps.reference_data.models import Currency, DocumentType

TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_api_tests_")

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
READ = ("generated_documents.read",)
GENERATE = ("generated_documents.generate",)
ALL_PERMS = READ + GENERATE


def soffice_available():
    return _find_soffice() is not None


def build_template_docx():
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class GeneratedDocumentAPITestCase(APITestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.company_a = make_company()
        self.company_b = make_company()
        self.alice = make_user(self.company_a, email="alice@example.com")
        self.bob = make_user(self.company_b, email="bob@example.com")
        self.document_type, _ = DocumentType.objects.get_or_create(
            code="invoice", defaults={"name": "Invoice"}
        )
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]

        self.template_a, self.template_version_a = self._make_template(self.company_a)
        self.template_b, self.template_version_b = self._make_template(self.company_b)
        self.client_a = self._make_client(self.company_a, "Acme A")
        self.client_b = self._make_client(self.company_b, "Acme B")
        self.doc_a = self._make_document(
            self.company_a, self.client_a, self.template_a, number="INV-A-001"
        )
        self.doc_b = self._make_document(
            self.company_b, self.client_b, self.template_b, number="INV-B-001"
        )

    # ------------------------------------------------------------------ setup

    def _make_template(self, company):
        template = DocumentTemplate.objects.create(
            company=company, document_type=self.document_type, name="Invoice Template"
        )
        template_version = TemplateVersion.objects.create(
            template=template,
            company=company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        template_version.file.save(
            "tpl.docx", ContentFile(build_template_docx()), save=True
        )
        return template, template_version

    def _make_client(self, company, legal_name):
        client = Client.objects.create(
            company=company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client, company=company, legal_name=legal_name, voen="1234567890"
        )
        return client

    def _make_document(self, company, client, template, number):
        return GeneratedDocument.objects.create(
            company=company,
            client=client,
            document_type=self.document_type,
            template=template,
            document_number=number,
        )

    def grant(self, user, *codes):
        role = Role.objects.create(
            name=f"Role {uuid.uuid4().hex[:8]}", company=None
        )
        for code in codes:
            role.permissions.add(Permission.objects.get(code=code))
        UserRole.objects.create(user=user, role=role)

    def auth(self, user):
        self.client.force_authenticate(user=user)
        return self.client

    def deauth(self):
        self.client.force_authenticate(user=None)

    def generate_docx_via_service(self, document):
        """Reuse the Step 4.4 service to produce a stored DOCX version."""
        version = generate_document(document, currency=self.currency)
        self.assertTrue(version.storage_path_docx)
        return version

    def add_pdf_via_service(self, version):
        generate_pdf(version)
        version.refresh_from_db()
        return version

    def response_body(self, response):
        if getattr(response, "streaming_content", None) is not None:
            return b"".join(response.streaming_content)
        return response.content

    # ------------------------------------------------------------------ urls

    def generate_url(self, document_id):
        return reverse("generated-document-generate", kwargs={"pk": document_id})

    def version_list_url(self, document_id):
        return reverse("generated-document-version-list", kwargs={"pk": document_id})

    def generate_pdf_url(self, version_id):
        return reverse(
            "generated-document-version-generate-pdf", kwargs={"version_id": version_id}
        )

    def docx_url(self, version_id):
        return reverse(
            "generated-document-version-docx", kwargs={"version_id": version_id}
        )

    def pdf_url(self, version_id):
        return reverse(
            "generated-document-version-pdf", kwargs={"version_id": version_id}
        )


class GenerationApiTests(GeneratedDocumentAPITestCase):
    def test_unauthenticated_generation_is_rejected(self):
        self.deauth()
        response = self.client.post(self.generate_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_generation_requires_permission(self):
        self.auth(self.alice)
        response = self.client.post(self.generate_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_authenticated_user_can_generate_document(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        response = self.client.post(
            self.generate_url(self.doc_a.id),
            {"currency_id": self.currency.id},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.json()
        self.assertEqual(data["version_number"], 1)
        self.assertEqual(data["generated_document_id"], self.doc_a.id)
        self.assertEqual(data["company_id"], self.company_a.id)
        self.assertEqual(data["template_version_id"], self.template_version_a.id)
        self.assertEqual(data["storage"], {"docx": True, "pdf": False})
        self.assertEqual(data["generated_by_id"], self.alice.id)
        self.assertTrue(GeneratedDocumentVersion.objects.filter(
            generated_document=self.doc_a, version_number=1
        ).exists())

    def test_generate_with_unknown_currency_returns_400(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        response = self.client.post(
            self.generate_url(self.doc_a.id),
            {"currency_id": 999999},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(GeneratedDocumentVersion.objects.filter(
            generated_document=self.doc_a
        ).exists())

    def test_regeneration_creates_new_versions_and_keeps_previous(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        first = self.client.post(
            self.generate_url(self.doc_a.id),
            {"currency_id": self.currency.id},
            format="json",
        )
        second = self.client.post(
            self.generate_url(self.doc_a.id),
            {"currency_id": self.currency.id},
            format="json",
        )
        self.assertEqual(first.status_code, status.HTTP_201_CREATED)
        self.assertEqual(second.status_code, status.HTTP_201_CREATED)
        self.assertEqual(first.json()["version_number"], 1)
        self.assertEqual(second.json()["version_number"], 2)

        listing = self.client.get(self.version_list_url(self.doc_a.id))
        self.assertEqual(listing.status_code, status.HTTP_200_OK)
        numbers = [row["version_number"] for row in listing.json()]
        self.assertEqual(numbers, [1, 2])
        # Version 1's DOCX is still downloadable.
        v1 = GeneratedDocumentVersion.objects.get(
            generated_document=self.doc_a, version_number=1
        )
        download = self.client.get(self.docx_url(v1.id))
        self.assertEqual(download.status_code, status.HTTP_200_OK)

    def test_generation_failure_returns_controlled_error(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        with mock.patch(
            "apps.generated_documents.api.views.generate_document",
            side_effect=GenerationError("boom"),
        ):
            response = self.client.post(self.generate_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("detail", response.json())
        self.assertNotIn("boom", response.json()["detail"])
        self.assertFalse(GeneratedDocumentVersion.objects.filter(
            generated_document=self.doc_a
        ).exists())

    def test_generation_failure_leaves_no_partial_success(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        # First generation succeeds...
        self.client.post(self.generate_url(self.doc_a.id))
        before = list(GeneratedDocumentVersion.objects.filter(
            generated_document=self.doc_a
        ).values_list("version_number", flat=True))
        # ...second fails; the existing version must remain intact.
        with mock.patch(
            "apps.generated_documents.api.views.generate_document",
            side_effect=GenerationError("boom"),
        ):
            response = self.client.post(self.generate_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            list(GeneratedDocumentVersion.objects.filter(
                generated_document=self.doc_a
            ).values_list("version_number", flat=True)),
            before,
        )


class AuthorizationTests(GeneratedDocumentAPITestCase):
    def test_read_without_generate_cannot_generate(self):
        self.grant(self.alice, *READ)
        self.auth(self.alice)
        response = self.client.post(self.generate_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_generate_without_read_cannot_download(self):
        self.grant(self.alice, *GENERATE)
        self.auth(self.alice)
        version = self.generate_docx_via_service(self.doc_a)
        response = self.client.get(self.docx_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)


class VersionListTests(GeneratedDocumentAPITestCase):
    def test_version_list_exposes_metadata_not_paths(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        version = self.generate_docx_via_service(self.doc_a)
        response = self.client.get(self.version_list_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = response.json()
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["id"], version.id)
        self.assertEqual(row["version_number"], 1)
        self.assertEqual(row["template_version_id"], self.template_version_a.id)
        self.assertEqual(row["storage"], {"docx": True, "pdf": False})
        self.assertIn("created_at", row)
        self.assertIn("generated_by_id", row)
        # Storage paths are internal and must never be exposed.
        self.assertNotIn("storage_path", "".join(row.keys()))
        self.assertNotIn(version.storage_path_docx, str(row))


class TenantIsolationTests(GeneratedDocumentAPITestCase):
    def setUp(self):
        super().setUp()
        self.grant(self.alice, *ALL_PERMS)
        self.grant(self.bob, *ALL_PERMS)
        self.auth(self.alice)
        self.doc_a_version = self.generate_docx_via_service(self.doc_a)

    def test_company_b_cannot_generate_company_a_document(self):
        self.auth(self.bob)
        response = self.client.post(self.generate_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_company_b_cannot_list_company_a_versions(self):
        self.auth(self.bob)
        response = self.client.get(self.version_list_url(self.doc_a.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_company_b_cannot_download_company_a_docx(self):
        self.auth(self.bob)
        response = self.client.get(self.docx_url(self.doc_a_version.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_company_b_cannot_download_company_a_pdf(self):
        if soffice_available():
            self.add_pdf_via_service(self.doc_a_version)
        self.auth(self.bob)
        response = self.client.get(self.pdf_url(self.doc_a_version.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_company_a_list_is_scoped_to_own_company(self):
        self.auth(self.alice)
        response = self.client.get(reverse("generated-document-list"))
        ids = {row["id"] for row in response.json()}
        self.assertEqual(ids, {self.doc_a.id})


class DocxDownloadTests(GeneratedDocumentAPITestCase):
    def setUp(self):
        super().setUp()
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        self.version = self.generate_docx_via_service(self.doc_a)

    def test_docx_download_returns_valid_docx(self):
        response = self.client.get(self.docx_url(self.version.id))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"].split(";")[0], DOCX_MIME)
        body = self.response_body(response)
        self.assertTrue(body)
        # A real DOCX is a ZIP archive (PK magic).
        self.assertTrue(body.startswith(b"PK"))

    def test_docx_download_has_safe_filename(self):
        response = self.client.get(self.docx_url(self.version.id))
        disposition = response["Content-Disposition"]
        self.assertIn("attachment", disposition)
        self.assertIn("INV-A-001-v1.docx", disposition)
        self.assertNotIn(self.version.storage_path_docx, disposition)
        self.assertNotIn("..", disposition)

    def test_docx_download_does_not_leak_storage_path(self):
        response = self.client.get(self.docx_url(self.version.id))
        body = self.response_body(response)
        self.assertNotIn(self.version.storage_path_docx.encode(), body)

    def test_docx_download_when_no_docx_returns_404(self):
        bare = GeneratedDocumentVersion.objects.create(
            generated_document=self.doc_a,
            company=self.company_a,
            version_number=99,
            template_version=self.template_version_a,
            currency=self.currency,
            generation_snapshot={},
            storage_path_docx=None,
        )
        response = self.client.get(self.docx_url(bare.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_docx_download_when_file_missing_on_storage_returns_404(self):
        orphan = GeneratedDocumentVersion.objects.create(
            generated_document=self.doc_a,
            company=self.company_a,
            version_number=98,
            template_version=self.template_version_a,
            currency=self.currency,
            generation_snapshot={},
            storage_path_docx="generated/company_0/doc_0/v98/gone.docx",
        )
        response = self.client.get(self.docx_url(orphan.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)


class PdfDownloadTests(GeneratedDocumentAPITestCase):
    def setUp(self):
        super().setUp()
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)

    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_pdf_download_returns_real_pdf(self):
        version = self.generate_docx_via_service(self.doc_a)
        self.add_pdf_via_service(version)

        response = self.client.get(self.pdf_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"].split(";")[0], "application/pdf")
        body = self.response_body(response)
        self.assertTrue(body.startswith(b"%PDF"))
        disposition = response["Content-Disposition"]
        self.assertIn("INV-A-001-v1.pdf", disposition)
        self.assertNotIn(version.storage_path_pdf, disposition)

    def test_pdf_download_when_no_pdf_returns_404(self):
        version = self.generate_docx_via_service(self.doc_a)
        response = self.client.get(self.pdf_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_unauthenticated_pdf_download_rejected(self):
        version = self.generate_docx_via_service(self.doc_a)
        self.deauth()
        response = self.client.get(self.pdf_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


class GeneratePdfApiTests(GeneratedDocumentAPITestCase):
    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_generate_pdf_endpoint_populates_pdf(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        version = self.generate_docx_via_service(self.doc_a)
        response = self.client.post(self.generate_pdf_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["id"], version.id)
        self.assertEqual(data["storage"], {"docx": True, "pdf": True})
        version.refresh_from_db()
        self.assertTrue(version.storage_path_pdf)

    def test_generate_pdf_requires_permission(self):
        self.auth(self.alice)
        version = self.generate_docx_via_service(self.doc_a)
        response = self.client.post(self.generate_pdf_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_generate_pdf_failure_is_controlled(self):
        self.grant(self.alice, *ALL_PERMS)
        self.auth(self.alice)
        version = self.generate_docx_via_service(self.doc_a)
        with mock.patch(
            "apps.generated_documents.api.views.generate_pdf",
            side_effect=PdfGenerationError("boom"),
        ):
            response = self.client.post(self.generate_pdf_url(version.id))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("detail", response.json())
        self.assertNotIn("boom", response.json()["detail"])
        version.refresh_from_db()
        self.assertFalse(version.storage_path_pdf)