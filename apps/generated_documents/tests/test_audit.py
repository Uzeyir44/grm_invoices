"""
Step 5.2 integration tests: generated-document operations (DOCX generation,
regeneration, PDF generation) produce audit events; failed generation /
failed PDF conversion produce none.
"""

import os
import shutil
import tempfile
import uuid
from io import BytesIO
from unittest import mock, skipUnless

from django.core.files.base import ContentFile
from django.test import override_settings
from django.urls import reverse
from docx import Document

from rest_framework import status
from rest_framework.test import APITestCase

from apps.audit_logs.models import AuditLog
from apps.authentication.models import Permission, Role, UserRole
from apps.authentication.tests.test_authentication import make_company, make_user
from apps.clients.models import Client, ClientLegalEntity
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import GeneratedDocument
from apps.generated_documents.services.context import GenerationContextError
from apps.generated_documents.services.generator import generate_document
from apps.generated_documents.services.pdf_generator import (
    PdfConversionUnavailableError,
    _find_soffice,
    generate_pdf,
)
from apps.reference_data.models import Currency, DocumentType

TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_audit_")

ALL_PERMS = ("generated_documents.read", "generated_documents.generate")


def soffice_available():
    return _find_soffice() is not None


def build_template_docx():
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class GeneratedDocumentAuditTestCase(APITestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.company = make_company()
        self.user = make_user(self.company, email="alice@example.com")
        self.other_company = make_company()
        self.other_user = make_user(self.other_company, email="bob@example.com")
        self.document_type = DocumentType.objects.create(code="invoice", name="Invoice")
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        self.template = DocumentTemplate.objects.create(
            company=self.company, document_type=self.document_type, name="Invoice Template"
        )
        self.template_version = TemplateVersion.objects.create(
            template=self.template,
            company=self.company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        self.template_version.file.save(
            "tpl.docx", ContentFile(build_template_docx()), save=True
        )
        self.customer = self._make_client(self.company)
        self.document = self._make_document(self.company, number="INV-001")

    def _make_client(self, company):
        client = Client.objects.create(
            company=company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client, company=company, legal_name="Acme Ltd", voen="1234567890"
        )
        return client

    def _make_document(self, company, number):
        return GeneratedDocument.objects.create(
            company=company,
            client=self.customer if company == self.company else None,
            document_type=self.document_type,
            template=self.template,
            document_number=number,
        )

    def grant(self, user, *codes):
        role = Role.objects.create(name=f"R{uuid.uuid4().hex[:8]}", company=None)
        for code in codes:
            role.permissions.add(Permission.objects.get(code=code))
        UserRole.objects.create(user=user, role=role)

    def auth(self, user):
        self.client.force_authenticate(user=user)

    def generate_url(self, document_id):
        return reverse("generated-document-generate", kwargs={"pk": document_id})


class GenerateAuditTests(GeneratedDocumentAuditTestCase):
    def test_service_generation_records_event(self):
        version = generate_document(self.document, currency=self.currency, generated_by=self.user)
        event = AuditLog.objects.filter(
            entity_type="generated_document", entity_id=self.document.id
        ).get()
        self.assertEqual(event.action, "generate")
        self.assertEqual(event.company_id, self.company.id)
        self.assertEqual(event.user_id, self.user.id)
        self.assertIsNone(event.previous_values)
        self.assertEqual(event.new_values["version_number"], version.version_number)
        self.assertEqual(event.new_values["document_number"], "INV-001")
        self.assertEqual(event.new_values["status"], "generated")

    def test_regeneration_records_a_new_event(self):
        generate_document(self.document, currency=self.currency, generated_by=self.user)
        generate_document(self.document, currency=self.currency, generated_by=self.user)
        events = AuditLog.objects.filter(
            entity_type="generated_document", entity_id=self.document.id
        ).order_by("created_at")
        self.assertEqual(events.count(), 2)
        self.assertEqual(events[0].new_values["version_number"], 1)
        self.assertEqual(events[1].new_values["version_number"], 2)

    def test_failed_generation_records_no_event(self):
        before = AuditLog.objects.count()
        doc = GeneratedDocument.objects.create(
            company=self.company,
            client=None,
            document_type=self.document_type,
            template=self.template,
            document_number="INV-NONE",
        )
        with self.assertRaises(GenerationContextError):
            generate_document(doc)
        self.assertEqual(AuditLog.objects.count(), before)

    def test_api_generation_records_event_with_user_and_metadata(self):
        self.grant(self.user, *ALL_PERMS)
        self.auth(self.user)
        response = self.client.post(
            self.generate_url(self.document.id),
            {"currency_id": self.currency.id},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        event = AuditLog.objects.filter(
            entity_type="generated_document", entity_id=self.document.id
        ).get()
        self.assertEqual(event.action, "generate")
        self.assertEqual(event.user_id, self.user.id)
        self.assertEqual(event.company_id, self.company.id)
        # Generation is audited inside the service layer (no HttpRequest), so
        # no request metadata is recorded for this event.
        self.assertIsNone(event.ip_address)

    def test_cross_company_generation_is_blocked_and_not_audited(self):
        self.grant(self.user, *ALL_PERMS)
        self.grant(self.other_user, *ALL_PERMS)
        before = AuditLog.objects.filter(company=self.other_company).count()
        self.auth(self.other_user)
        response = self.client.post(
            self.generate_url(self.document.id), {}, format="json"
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(
            AuditLog.objects.filter(company=self.other_company).count(), before
        )


class GeneratePdfAuditTests(GeneratedDocumentAuditTestCase):
    def _docx_version(self):
        return generate_document(
            self.document, currency=self.currency, generated_by=self.user
        )

    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_successful_pdf_records_event(self):
        version = self._docx_version()
        generate_pdf(version, generated_by=self.user)
        event = AuditLog.objects.filter(
            entity_type="generated_document_version", entity_id=version.id
        ).get()
        self.assertEqual(event.action, "generate_pdf")
        self.assertEqual(event.company_id, self.company.id)
        self.assertEqual(event.user_id, self.user.id)
        self.assertEqual(event.new_values["generated_document_id"], self.document.id)
        self.assertEqual(event.new_values["version_number"], version.version_number)

    def test_failed_pdf_records_no_event(self):
        version = self._docx_version()

        def exploding(docx_bytes):
            raise PdfConversionUnavailableError("synthetic failure")

        with self.assertRaises(PdfConversionUnavailableError):
            generate_pdf(version, converter=exploding, generated_by=self.user)
        self.assertFalse(
            AuditLog.objects.filter(
                entity_type="generated_document_version", entity_id=version.id
            ).exists()
        )

    def test_failed_api_pdf_records_no_event(self):
        self.grant(self.user, *ALL_PERMS)
        self.auth(self.user)
        version = self._docx_version()
        with mock.patch(
            "apps.generated_documents.api.views.generate_pdf",
            side_effect=PdfConversionUnavailableError("unavailable"),
        ):
            response = self.client.post(
                reverse(
                    "generated-document-version-generate-pdf",
                    kwargs={"version_id": version.id},
                ),
                {},
                format="json",
            )
        self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        # A failed conversion through the API must not produce an event.
        self.assertFalse(
            AuditLog.objects.filter(
                entity_type="generated_document_version", entity_id=version.id
            ).exists()
        )