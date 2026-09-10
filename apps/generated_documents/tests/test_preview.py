"""
Non-persistent document preview API tests.

``POST /api/generated-documents/preview/`` renders the current form data into a
temporary PDF using the real generation pipeline, but must persist nothing: no
document, version, line items, audit event, consumed number, or stored file.
The LibreOffice converter is stubbed so these tests never need the runtime.
"""

import os
import shutil
import tempfile
import uuid
import zipfile
from io import BytesIO
from unittest import mock

from django.core.files.base import ContentFile
from django.test import override_settings
from django.urls import reverse
from docx import Document

from rest_framework import status
from rest_framework.test import APITestCase

from apps.audit_logs.models import AuditLog
from apps.authentication.models import Permission, Role, UserRole
from apps.authentication.tests.test_authentication import make_company, make_user
from apps.clients.models import Client, ClientLegalEntity
from apps.company_document_counters.models import CompanyDocumentCounter
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import (
    DocumentLineItem,
    GeneratedDocument,
    GeneratedDocumentVersion,
)
from apps.generated_documents.services.pdf_generator import (
    PdfConversionUnavailableError,
)
from apps.reference_data.models import Currency, DocumentType

TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_preview_tests_")

GENERATE = ("generated_documents.generate",)

CONVERTER_PATH = (
    "apps.generated_documents.services.preview.convert_docx_bytes_to_pdf"
)
FAKE_PDF = b"%PDF-1.4 fake preview"


def build_template_docx():
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


def extract_docx_text(docx_bytes):
    with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    # Strip tags so the rendered paragraph text can be asserted directly.
    import re

    return re.sub(r"<[^>]+>", "", xml)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class PreviewApiTests(APITestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        shutil.rmtree(os.path.join(TEMP_MEDIA, "generated"), ignore_errors=True)
        self.company_a = make_company()
        self.company_b = make_company()
        self.alice = make_user(self.company_a, email="alice@example.com")
        self.bob = make_user(self.company_b, email="bob@example.com")
        self.document_type, _ = DocumentType.objects.get_or_create(
            code="invoice", defaults={"name": "Invoice"}
        )
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        self.template_a, self.template_version_a = self._make_template(
            self.company_a, self.document_type, build_template_docx()
        )
        self.client_a = self._make_client(self.company_a, "Acme A")
        self.client_b = self._make_client(self.company_b, "Acme B")

    def _make_template(self, company, document_type, file_bytes):
        template = DocumentTemplate.objects.create(
            company=company, document_type=document_type, name="Invoice Template"
        )
        template_version = TemplateVersion.objects.create(
            template=template,
            company=company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        template_version.file.save(
            "tpl.docx", ContentFile(file_bytes), save=True
        )
        return template, template_version

    def _make_client(self, company, legal_name):
        client = Client.objects.create(
            company=company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client, company=company, legal_name=legal_name, voen="1234567890"
        )
        return client

    def grant(self, user, *codes):
        role = Role.objects.create(
            name=f"Role {uuid.uuid4().hex[:8]}", company=None
        )
        for code in codes:
            role.permissions.add(Permission.objects.get(code=code))
        UserRole.objects.create(user=user, role=role)

    def url(self):
        return reverse("generated-document-preview")

    def payload(self, **overrides):
        body = {
            "client_id": self.client_a.id,
            "document_type_id": self.document_type.id,
            "currency_id": self.currency.id,
            "items": [
                {
                    "description": "Consulting hours",
                    "quantity": "10.0000",
                    "unit": "hour",
                    "unit_price": "100.0000",
                    "discount_percent": "0.00",
                    "vat_rate_percent": "18.00",
                }
            ],
        }
        body.update(overrides)
        return body

    def response_body(self, response):
        if getattr(response, "streaming_content", None) is not None:
            return b"".join(response.streaming_content)
        return response.content

    def generated_storage_exists(self):
        path = os.path.join(TEMP_MEDIA, "generated")
        if not os.path.isdir(path):
            return False
        for _, _, files in os.walk(path):
            if files:
                return True
        return False

    def assert_nothing_persisted(self):
        self.assertFalse(GeneratedDocument.objects.exists())
        self.assertFalse(GeneratedDocumentVersion.objects.exists())
        self.assertFalse(DocumentLineItem.objects.exists())
        self.assertFalse(AuditLog.objects.exists())
        self.assertFalse(CompanyDocumentCounter.objects.exists())
        self.assertFalse(self.generated_storage_exists())

    # ------------------------------------------------------------ success

    def test_success_returns_pdf_and_persists_nothing(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(self.url(), self.payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"].split(";")[0], "application/pdf"
        )
        body = self.response_body(response)
        self.assertTrue(body.startswith(b"%PDF"))
        self.assert_nothing_persisted()

    def test_preview_uses_provisional_number(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        captured = {}

        def capturing_converter(docx_bytes):
            captured["docx"] = docx_bytes
            return FAKE_PDF

        with mock.patch(CONVERTER_PATH, side_effect=capturing_converter):
            response = self.client.post(self.url(), self.payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        text = extract_docx_text(captured["docx"])
        self.assertIn("PREVIEW", text)
        self.assertIn("Invoice: PREVIEW", text)
        # A real generated document would have a numeric number.
        self.assertNotIn("Invoice: 1", text)
        self.assert_nothing_persisted()

    def test_client_supplied_company_id_is_ignored(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(),
                self.payload(company_id=self.company_b.id),
                format="json",
            )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assert_nothing_persisted()

    # ------------------------------------------------------------ security

    def test_unauthenticated_is_rejected(self):
        response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assert_nothing_persisted()

    def test_preview_requires_permission(self):
        self.client.force_authenticate(user=self.alice)
        response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assert_nothing_persisted()

    def test_cross_tenant_client_resolves_to_404(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(), self.payload(client_id=self.client_b.id), format="json"
            )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assert_nothing_persisted()

    # ------------------------------------------------------------ validation

    def test_missing_fields_return_400(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        response = self.client.post(self.url(), {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assert_nothing_persisted()

    def test_unknown_currency_returns_400(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(), self.payload(currency_id=999999), format="json"
            )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assert_nothing_persisted()

    def test_no_active_template_returns_400(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        other_type = DocumentType.objects.create(code="contract", name="Contract")
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(),
                self.payload(document_type_id=other_type.id),
                format="json",
            )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assert_nothing_persisted()

    def test_invalid_placeholder_returns_400(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        bad_type = DocumentType.objects.create(code="bad", name="Bad")
        bad_docx = Document()
        bad_docx.add_paragraph("Bad: {{ some_random_variable }}")
        out = BytesIO()
        bad_docx.save(out)
        self._make_template(self.company_a, bad_type, out.getvalue())
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(),
                self.payload(document_type_id=bad_type.id),
                format="json",
            )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assert_nothing_persisted()

    # ------------------------------------------------------------ PDF errors

    def test_pdf_runtime_unavailable_returns_503(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(
            CONVERTER_PATH, side_effect=PdfConversionUnavailableError("no runtime")
        ):
            response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(
            response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE
        )
        self.assert_nothing_persisted()

    def test_invalid_pdf_output_returns_400(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=b"not a pdf"):
            response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assert_nothing_persisted()
