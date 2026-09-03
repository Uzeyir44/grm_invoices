"""
Tests for document creation + automatic numbering (service + API).

Covers: basic creation (document/initial version/line items/audit), company
scoping of the number sequence, per-company independence, tenant isolation,
RBAC, rollback on version/line-item/audit failure, and integration with the
existing DOCX/PDF generation pipeline. Includes a concurrency test proving the
company-scoped counter never yields duplicate numbers.
"""

import shutil
import tempfile
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest import mock, skipUnless

from django.core.files.base import ContentFile
from django.db import IntegrityError, transaction
from django.test import TransactionTestCase, override_settings
from django.urls import reverse
from docx import Document

from rest_framework import status
from rest_framework.test import APITestCase

from apps.audit_logs.models import AuditLog
from apps.audit_logs.services import AuditLogError
from apps.authentication.models import Permission, Role, UserRole
from apps.authentication.tests.test_authentication import make_company, make_user
from apps.clients.models import Client, ClientLegalEntity
from apps.company_document_counters.models import CompanyDocumentCounter
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import (
    DocumentLineItem,
    GeneratedDocument,
    GeneratedDocumentVersion,
)
from apps.generated_documents.services.document_creation import (
    DocumentCreationError,
    create_generated_document,
)
from apps.generated_documents.services.pdf_generator import _find_soffice
from apps.reference_data.models import Currency, DocumentType

TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_doc_create_")

PERMS = ("generated_documents.read", "generated_documents.generate")


def build_invoice_template():
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    table = document.add_table(rows=0, cols=3)
    header = table.add_row().cells
    header[0].text = "Description"
    header[1].text = "Qty"
    header[2].text = "Total"
    proto = table.add_row().cells
    proto[0].text = "{{#each items}}"
    proto[1].text = "{{ item.description }}"
    proto[2].text = "{{/each}}"
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


def soffice_available():
    return _find_soffice() is not None


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class DocumentCreationTestCase(APITestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.company_a = make_company()
        self.company_b = make_company()
        self.alice = make_user(self.company_a, email="alice@example.com")
        self.carol = make_user(self.company_b, email="carol@example.com")
        self.document_type = DocumentType.objects.create(code="invoice", name="Invoice")
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        self.template_a, self.template_version_a = self._make_template(self.company_a)
        self.template_b, self.template_version_b = self._make_template(self.company_b)
        self.client_a = self._make_client(self.company_a, "Acme A")
        self.client_b = self._make_client(self.company_b, "Acme B")

    def _make_template(self, company):
        template = DocumentTemplate.objects.create(
            company=company,
            document_type=self.document_type,
            name="Invoice Template",
        )
        template_version = TemplateVersion.objects.create(
            template=template,
            company=company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        template_version.file.save(
            "tpl.docx", ContentFile(build_invoice_template()), save=True
        )
        return template, template_version

    def _make_client(self, company, legal_name):
        client = Client.objects.create(
            company=company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client, company=company, legal_name=legal_name, voen="1234567890"
        )
        return client

    def grant(self, user, *codes):
        role = Role.objects.create(
            name=f"Role {uuid.uuid4().hex[:8]}", company=None
        )
        for code in codes:
            role.permissions.add(Permission.objects.get(code=code))
        UserRole.objects.create(user=user, role=role)

    def auth(self, user):
        self.client.force_authenticate(user=user)

    def create_url(self):
        return reverse("generated-document-list")

    def payload(self, **overrides):
        payload = {
            "client_id": self.client_a.id,
            "document_type_id": self.document_type.id,
            "currency_id": self.currency.id,
            "items": [
                {
                    "description": "Consulting",
                    "quantity": "2",
                    "unit": "h",
                    "unit_price": "100",
                    "discount_percent": "0",
                    "vat_rate_percent": "18",
                }
            ],
        }
        payload.update(overrides)
        return payload

    def api_create(self, payload):
        return self.client.post(self.create_url(), payload, format="json")


class BasicCreationTests(DocumentCreationTestCase):
    def test_create_creates_document_version_items_and_audit(self):
        self.grant(self.alice, *PERMS)
        self.auth(self.alice)
        response = self.api_create(self.payload())
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.json()
        self.assertEqual(data["document_number"], "1")
        self.assertEqual(data["status"], "draft")
        self.assertEqual(data["company_id"], self.company_a.id)
        self.assertEqual(data["client_id"], self.client_a.id)
        self.assertEqual(data["template_id"], self.template_a.id)
        self.assertEqual(data["document_type"], self.document_type.id)
        self.assertEqual(data["created_by_id"], self.alice.id)

        document = GeneratedDocument.objects.get(pk=data["id"])
        self.assertEqual(document.document_number, "1")
        self.assertEqual(document.versions.count(), 1)
        version = document.versions.get()
        self.assertEqual(version.version_number, 1)
        self.assertEqual(version.template_version_id, self.template_version_a.id)
        self.assertEqual(version.currency_id, self.currency.id)
        self.assertFalse(version.storage_path_docx)
        self.assertEqual(version.generated_by_id, self.alice.id)

        item = version.line_items.get()
        self.assertEqual(item.description, "Consulting")
        self.assertEqual(item.quantity, Decimal("2.0000"))
        self.assertEqual(item.line_number, 1)
        self.assertEqual(item.subtotal, Decimal("200.0000"))
        self.assertEqual(item.vat_amount, Decimal("36.0000"))
        self.assertEqual(item.total, Decimal("236.0000"))

        events = AuditLog.objects.filter(
            entity_type="generated_document", entity_id=document.id
        )
        self.assertEqual(events.count(), 1)
        event = events.get()
        self.assertEqual(event.action, "create")
        self.assertEqual(event.company_id, self.company_a.id)
        self.assertEqual(event.user_id, self.alice.id)
        self.assertEqual(event.new_values["document_number"], "1")
        self.assertEqual(event.new_values["version_number"], 1)
        self.assertEqual(event.ip_address, "127.0.0.1")

    def test_subsequent_documents_increment_number(self):
        self.grant(self.alice, *PERMS)
        self.auth(self.alice)
        first = self.api_create(self.payload())
        second = self.api_create(self.payload())
        self.assertEqual(first.json()["document_number"], "1")
        self.assertEqual(second.json()["document_number"], "2")

    def test_document_type_and_currency_are_verified(self):
        self.grant(self.alice, *PERMS)
        self.auth(self.alice)
        unknown_type = self.api_create(self.payload(document_type_id=999999))
        self.assertEqual(unknown_type.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(
            GeneratedDocument.objects.filter(company=self.company_a).exists()
        )
        missing_currency = self.api_create(
            self.payload(currency_id=None, items=[])
        )
        self.assertEqual(missing_currency.status_code, status.HTTP_400_BAD_REQUEST)
        # Failed creations rolled back the counter too: next number is still 1.
        ok = self.api_create(self.payload())
        self.assertEqual(ok.status_code, status.HTTP_201_CREATED)
        self.assertEqual(ok.json()["document_number"], "1")


class NumberingScopeTests(DocumentCreationTestCase):
    def test_independent_counters_per_company(self):
        self.grant(self.alice, *PERMS)
        self.grant(self.carol, *PERMS)
        self.auth(self.alice)
        self.assertEqual(
            self.api_create(self.payload()).json()["document_number"], "1"
        )
        self.auth(self.carol)
        carol_payload = self.payload()
        carol_payload.update(
            client_id=self.client_b.id,
        )
        self.assertEqual(
            self.api_create(carol_payload).json()["document_number"], "1"
        )
        self.auth(self.alice)
        self.assertEqual(
            self.api_create(self.payload()).json()["document_number"], "2"
        )
        # Counter rows are scoped per company.
        self.assertEqual(
            CompanyDocumentCounter.objects.get(company=self.company_a).last_number,
            2,
        )
        self.assertEqual(
            CompanyDocumentCounter.objects.get(company=self.company_b).last_number,
            1,
        )


class CreationSecurityTests(DocumentCreationTestCase):
    def test_unauthenticated_request_is_denied(self):
        response = self.client.post(self.create_url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_requires_generate_permission(self):
        self.auth(self.alice)
        response = self.api_create(self.payload())
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_cannot_use_another_companys_client(self):
        self.grant(self.carol, *PERMS)
        self.auth(self.carol)
        payload = self.payload(client_id=self.client_a.id)
        response = self.api_create(payload)
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertFalse(
            GeneratedDocument.objects.filter(company=self.company_b).exists()
        )

    def test_client_supplied_company_id_is_ignored(self):
        self.grant(self.alice, *PERMS)
        self.auth(self.alice)
        payload = self.payload(company_id=self.company_b.id)
        response = self.api_create(payload)
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["company_id"], self.company_a.id)
        self.assertFalse(
            GeneratedDocument.objects.filter(company=self.company_b).exists()
        )


class RollbackTests(DocumentCreationTestCase):
    def _call_service(self, **kwargs):
        return create_generated_document(
            company=self.company_a,
            user=self.alice,
            client=self.client_a,
            document_type_id=self.document_type.id,
            currency_id=self.currency.id,
            items=[
                {
                    "description": "Consulting",
                    "quantity": Decimal("2"),
                    "unit": "h",
                    "unit_price": Decimal("100"),
                    "discount_percent": Decimal("0"),
                    "vat_rate_percent": Decimal("18"),
                }
            ],
            **kwargs,
        )

    def test_failure_during_version_creation_rolls_back_everything(self):
        with mock.patch.object(
            GeneratedDocumentVersion.objects,
            "create",
            side_effect=IntegrityError("boom"),
        ):
            with self.assertRaises(IntegrityError):
                self._call_service()
        self.assertEqual(GeneratedDocument.objects.filter(company=self.company_a).count(), 0)
        self.assertEqual(
            AuditLog.objects.filter(entity_type="generated_document").count(), 0
        )
        # The counter allocation also rolled back.
        ok = self._call_service()
        self.assertEqual(ok.document_number, "1")

    def test_failure_during_line_item_creation_rolls_back_everything(self):
        with mock.patch.object(
            DocumentLineItem.objects,
            "create",
            side_effect=IntegrityError("boom"),
        ):
            with self.assertRaises(IntegrityError):
                self._call_service()
        self.assertEqual(
            GeneratedDocument.objects.filter(company=self.company_a).count(), 0
        )
        self.assertEqual(
            GeneratedDocumentVersion.objects.filter(
                generated_document__company=self.company_a
            ).count(),
            0,
        )

    def test_failure_during_audit_creation_rolls_back_everything(self):
        with mock.patch(
            "apps.generated_documents.services.document_creation.create_audit_log",
            side_effect=AuditLogError("boom"),
        ):
            with self.assertRaises(AuditLogError):
                self._call_service()
        self.assertEqual(
            GeneratedDocument.objects.filter(company=self.company_a).count(), 0
        )
        self.assertEqual(
            GeneratedDocumentVersion.objects.filter(
                generated_document__company=self.company_a
            ).count(),
            0,
        )
        self.assertEqual(AuditLog.objects.count(), 0)
        self.assertEqual(
            CompanyDocumentCounter.objects.filter(company=self.company_a).count(), 0
        )


class GenerationIntegrationTests(DocumentCreationTestCase):
    def setUp(self):
        super().setUp()
        self.grant(self.alice, *PERMS)
        self.auth(self.alice)
        self.document_id = self.api_create(self.payload()).json()["id"]

    def _docx_text(self, version):
        import zipfile

        from django.core.files.storage import default_storage

        with default_storage.open(version.storage_path_docx, "rb") as fobj:
            xml = zipfile.ZipFile(fobj).read("word/document.xml")
        from lxml import etree

        root = etree.fromstring(xml)
        return "".join(
            t.text or "" for t in root.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
            )
        )

    def test_created_document_can_flow_into_generation_pipeline(self):
        response = self.client.post(
            reverse("generated-document-generate", kwargs={"pk": self.document_id}),
            {},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.json()
        # Creation created v1; generation produces v2 with the rendered DOCX.
        self.assertEqual(data["version_number"], 2)
        self.assertEqual(data["storage"], {"docx": True, "pdf": False})

        version = GeneratedDocumentVersion.objects.get(pk=data["id"])
        text = self._docx_text(version)
        self.assertIn("Consulting", text)
        self.assertIn("Invoice: 1", text)  # placeholder text came from the number
        self.assertNotIn("{{", text)

    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_created_document_pdf_still_works(self):
        generated = self.client.post(
            reverse("generated-document-generate", kwargs={"pk": self.document_id}),
            {},
            format="json",
        ).json()
        response = self.client.post(
            reverse(
                "generated-document-version-generate-pdf",
                kwargs={"version_id": generated["id"]},
            ),
            {},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["storage"], {"docx": True, "pdf": True})


class ConcurrencyTests(TransactionTestCase):
    """Concurrent creation cannot produce duplicate (company, number)."""

    def setUp(self):
        self.company = make_company()
        self.user = make_user(self.company, email="bob@example.com")
        self.document_type = DocumentType.objects.create(code="invoice", name="Invoice")
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        template = DocumentTemplate.objects.create(
            company=self.company,
            document_type=self.document_type,
            name="Invoice Template",
        )
        template_version = TemplateVersion.objects.create(
            template=template,
            company=self.company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        template_version.file.save(
            "tpl.docx", ContentFile(build_invoice_template()), save=True
        )
        client = Client.objects.create(
            company=self.company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client,
            company=self.company,
            legal_name="Acme",
            voen="1234567890",
        )
        self.client_id = client.id

    def test_simultaneous_creation_allocates_unique_numbers(self):
        def create(_):
            try:
                create_generated_document(
                    company=self.company,
                    user=self.user,
                    client=Client.objects.get(pk=self.client_id),
                    document_type_id=self.document_type.id,
                    currency_id=self.currency.id,
                    items=[
                        {
                            "description": "Item",
                            "quantity": Decimal("1"),
                            "unit_price": Decimal("10"),
                            "vat_rate_percent": Decimal("18"),
                        }
                    ],
                )
            finally:
                # Each worker opens its own DB connection; close it so the
                # test runner can clean up afterwards.
                from django.db import connection

                connection.close()

        with ThreadPoolExecutor(max_workers=8) as executor:
            list(executor.map(create, range(8)))

        docs = GeneratedDocument.objects.filter(company=self.company)
        self.assertEqual(docs.count(), 8)
        numbers = list(docs.values_list("document_number", flat=True))
        self.assertEqual(len(numbers), len(set(numbers)))
        # The DB unique constraint was satisfied as well.
        self.assertEqual(len(docs.values("document_number").distinct()), 8)
        self.assertEqual(
            AuditLog.objects.filter(
                company=self.company, action="create"
            ).count(),
            8,
        )