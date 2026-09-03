"""
Step 4.5 focused tests: DOCX -> PDF conversion of an already-generated
``GeneratedDocumentVersion`` (apps/generated_documents/services/pdf_generator.py).

Deliberately does not rewrite the Step 4.4 suite; it only exercises the PDF
layer on top of the existing ``generate_document`` pipeline.
"""

import os
import shutil
import tempfile
from io import BytesIO
from unittest import skipUnless

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import DatabaseError, transaction
from django.test import TestCase, override_settings
from docx import Document

from apps.authentication.tests.test_authentication import make_company
from apps.clients.models import Client, ClientLegalEntity
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import GeneratedDocument, GeneratedDocumentVersion
from apps.generated_documents.services.generator import generate_document
from apps.generated_documents.services.pdf_generator import (
    PdfConversionUnavailableError,
    PdfGenerationError,
    _find_soffice,
    generate_pdf,
)
from apps.reference_data.models import Currency, DocumentType

TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_pdf_tests_")


def build_template_docx():
    """Small real DOCX fixture: simple placeholders (no line items needed)."""
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


def soffice_available():
    return _find_soffice() is not None


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class PdfGenerationTestCase(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.company = make_company()
        self.document_type = DocumentType.objects.create(code="invoice", name="Invoice")
        self.template = DocumentTemplate.objects.create(
            company=self.company,
            document_type=self.document_type,
            name="Invoice Template",
        )
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        self.client = Client.objects.create(
            company=self.company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=self.client,
            company=self.company,
            legal_name="Acme Ltd",
            voen="1234567890",
        )
        self.template_version = TemplateVersion.objects.create(
            template=self.template,
            company=self.company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        self.template_version.file.save(
            "tpl.docx", ContentFile(build_template_docx()), save=True
        )
        self.document = GeneratedDocument.objects.create(
            company=self.company,
            client=self.client,
            document_type=self.document_type,
            template=self.template,
            document_number="INV-001",
        )

    def make_generated_version(self):
        """Reuse the Step 4.4 pipeline to produce a version with a stored DOCX."""
        version = generate_document(self.document, currency=self.currency)
        self.assertTrue(version.storage_path_docx)
        return version

    def read_stored(self, path):
        with default_storage.open(path, "rb") as fh:
            return fh.read()

    def stored_file_exists(self, path):
        return bool(path) and default_storage.exists(path)


class SuccessfulConversionTests(PdfGenerationTestCase):
    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_successful_conversion_produces_real_pdf(self):
        version = self.make_generated_version()
        docx_path = version.storage_path_docx
        docx_before = self.read_stored(docx_path)

        result = generate_pdf(version)

        self.assertIs(result, version)
        self.assertTrue(version.storage_path_pdf)
        self.assertTrue(self.stored_file_exists(version.storage_path_pdf))
        pdf_bytes = self.read_stored(version.storage_path_pdf)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"), "file must be a real PDF")
        self.assertGreater(len(pdf_bytes), 1000)
        # Same deterministic per-version directory as the DOCX (Step 4.4 layout).
        self.assertEqual(
            os.path.dirname(version.storage_path_pdf), os.path.dirname(docx_path)
        )
        # The DOCX (and its recorded path) must remain untouched.
        self.assertEqual(version.storage_path_docx, docx_path)
        self.assertEqual(self.read_stored(docx_path), docx_before)
        self.assertEqual(
            self.document.versions.count(),
            1,
            "PDF generation must not create a new GeneratedDocumentVersion",
        )

    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_repeated_conversion_replaces_same_pdf(self):
        version = self.make_generated_version()
        first = generate_pdf(version).storage_path_pdf
        second = generate_pdf(version).storage_path_pdf
        self.assertEqual(first, second)
        pdf_dir = os.path.join(TEMP_MEDIA, os.path.dirname(first))
        pdf_files = [n for n in os.listdir(pdf_dir) if n.endswith(".pdf")]
        self.assertEqual(len(pdf_files), 1, "no uncontrolled duplicate PDFs")

    @skipUnless(soffice_available(), "LibreOffice (soffice) not available")
    def test_pdf_generation_does_not_modify_source_template(self):
        template_path = self.template_version.file.name
        template_before = self.read_stored(template_path)

        version = self.make_generated_version()
        generate_pdf(version)

        self.assertEqual(self.read_stored(template_path), template_before)
        version.refresh_from_db()
        self.assertEqual(version.template_version_id, self.template_version.id)


class FailureBehaviorTests(PdfGenerationTestCase):
    def _exploding_converter(self, docx_bytes):
        raise PdfConversionUnavailableError("synthetic conversion failure")

    def test_conversion_failure_leaves_docx_intact_and_no_pdf(self):
        version = self.make_generated_version()
        docx_path = version.storage_path_docx
        docx_before = self.read_stored(docx_path)

        with self.assertRaises(PdfConversionUnavailableError):
            generate_pdf(version, converter=self._exploding_converter)

        version.refresh_from_db()
        self.assertIsNone(version.storage_path_pdf)
        self.assertEqual(version.storage_path_docx, docx_path)
        self.assertEqual(self.read_stored(docx_path), docx_before)
        self.assertEqual(self.document.versions.count(), 1)
        self.assertFalse(
            self.stored_file_exists(docx_path[: -len(".docx")] + ".pdf")
        )

    def test_invalid_pdf_bytes_are_rejected_and_nothing_stored(self):
        version = self.make_generated_version()

        with self.assertRaises(PdfGenerationError):
            generate_pdf(version, converter=lambda data: b"not a real pdf")

        version.refresh_from_db()
        self.assertIsNone(version.storage_path_pdf)
        self.assertFalse(
            self.stored_file_exists(
                version.storage_path_docx[: -len(".docx")] + ".pdf"
            )
        )

    def test_version_without_stored_docx_raises(self):
        version = GeneratedDocumentVersion.objects.create(
            generated_document=self.document,
            company=self.company,
            version_number=1,
            template_version=self.template_version,
            currency=self.currency,
            generation_snapshot={},
        )
        with self.assertRaises(PdfGenerationError):
            generate_pdf(version)
        version.refresh_from_db()
        self.assertIsNone(version.storage_path_pdf)


class ImmutabilityGuardTests(PdfGenerationTestCase):
    def test_other_columns_still_immutable_after_pdf(self):
        version = self.make_generated_version()

        with self.assertRaises((DatabaseError,)):
            with transaction.atomic():
                GeneratedDocumentVersion.objects.filter(pk=version.pk).update(
                    version_number=99
                )

        version.refresh_from_db()
        self.assertEqual(version.version_number, 1)
        self.assertIsNone(version.storage_path_pdf)

    def test_pdf_path_update_is_allowed_by_trigger(self):
        version = self.make_generated_version()

        # Direct UPDATE of ONLY storage_path_pdf must be allowed (what the
        # pdf generator does); the version snapshot columns stay immutable.
        GeneratedDocumentVersion.objects.filter(pk=version.pk).update(
            storage_path_pdf="generated/company_0/doc_0/v1/p.pdf"
        )
        version.refresh_from_db()
        self.assertEqual(
            version.storage_path_pdf, "generated/company_0/doc_0/v1/p.pdf"
        )
        self.assertEqual(version.version_number, 1)