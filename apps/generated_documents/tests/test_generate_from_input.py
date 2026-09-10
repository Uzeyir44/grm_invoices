"""
Atomic generate-from-input API tests.

``POST /api/generated-documents/generate-from-input/`` must create the
document, its version 1 (with DOCX + PDF), line items and audit records in ONE
transaction — and persist nothing (rows, numbers, or files) when any step
fails. The LibreOffice converter is stubbed so these tests never need the
system runtime.
"""

import os
import shutil
import tempfile
import uuid
from io import BytesIO
from unittest import mock

from django.core.files.base import ContentFile
from django.test import override_settings
from django.urls import reverse
from docx import Document

from rest_framework import status
from rest_framework.test import APITestCase

from apps.audit_logs.models import AuditLog
from apps.authentication.models import Permission, Role, UserRole
from apps.authentication.tests.test_authentication import make_company, make_user
from apps.clients.models import Client, ClientLegalEntity
from apps.company_document_counters.models import CompanyDocumentCounter
from apps.document_templates.models import DocumentTemplate, TemplateVersion
from apps.generated_documents.models import (
    DocumentLineItem,
    GeneratedDocument,
    GeneratedDocumentVersion,
)
from apps.generated_documents.services.pdf_generator import PdfGenerationError
from apps.reference_data.models import Currency, DocumentType

TEMP_MEDIA = tempfile.mkdtemp(prefix="grm_from_input_tests_")

GENERATE = ("generated_documents.generate",)

CONVERTER_PATH = (
    "apps.generated_documents.services.generate_from_input."
    "convert_docx_bytes_to_pdf"
)
FAKE_PDF = b"%PDF-1.4 fake"


def build_template_docx():
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class GenerateFromInputApiTests(APITestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        # TEMP_MEDIA is shared for the whole class run: clear previously
        # generated files so each test's storage assertions are hermetic.
        # (Template fixtures live under templates/, never under generated/.)
        shutil.rmtree(os.path.join(TEMP_MEDIA, "generated"), ignore_errors=True)
        self.company_a = make_company()
        self.company_b = make_company()
        self.alice = make_user(self.company_a, email="alice@example.com")
        self.bob = make_user(self.company_b, email="bob@example.com")
        self.document_type, _ = DocumentType.objects.get_or_create(
            code="invoice", defaults={"name": "Invoice"}
        )
        self.currency = Currency.objects.get_or_create(code="AZN", name="Manat")[0]
        self.template_a, self.template_version_a = self._make_template(self.company_a)
        self.client_a = self._make_client(self.company_a, "Acme A")
        self.client_b = self._make_client(self.company_b, "Acme B")

    def _make_template(self, company):
        template = DocumentTemplate.objects.create(
            company=company, document_type=self.document_type, name="Invoice Template"
        )
        template_version = TemplateVersion.objects.create(
            template=template,
            company=company,
            version_number=1,
            storage_path="v1",
            is_active=True,
        )
        template_version.file.save(
            "tpl.docx", ContentFile(build_template_docx()), save=True
        )
        return template, template_version

    def _make_client(self, company, legal_name):
        client = Client.objects.create(
            company=company, client_type=Client.ClientType.LEGAL_ENTITY
        )
        ClientLegalEntity.objects.create(
            client=client, company=company, legal_name=legal_name, voen="1234567890"
        )
        return client

    def grant(self, user, *codes):
        role = Role.objects.create(
            name=f"Role {uuid.uuid4().hex[:8]}", company=None
        )
        for code in codes:
            role.permissions.add(Permission.objects.get(code=code))
        UserRole.objects.create(user=user, role=role)

    def url(self):
        return reverse("generated-document-generate-from-input")

    def payload(self, **overrides):
        body = {
            "client_id": self.client_a.id,
            "document_type_id": self.document_type.id,
            "currency_id": self.currency.id,
            "items": [
                {
                    "description": "Consulting hours",
                    "quantity": "10.0000",
                    "unit": "hour",
                    "unit_price": "100.0000",
                    "discount_percent": "0.00",
                    "vat_rate_percent": "18.00",
                }
            ],
        }
        body.update(overrides)
        return body

    def generated_storage_exists(self):
        path = os.path.join(TEMP_MEDIA, "generated")
        if not os.path.isdir(path):
            return False
        for _, _, files in os.walk(path):
            if files:
                return True
        return False

    def counter_value(self, company):
        row = CompanyDocumentCounter.objects.filter(company=company).first()
        return None if row is None else row.last_number

    # ------------------------------------------------------------ success

    def test_success_creates_generated_document_with_docx_and_pdf(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(self.url(), self.payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.json()
        self.assertEqual(data["status"], "generated")
        self.assertEqual(data["company_id"], self.company_a.id)

        document = GeneratedDocument.objects.get(pk=data["id"])
        self.assertEqual(document.status, GeneratedDocument.Status.GENERATED)
        self.assertEqual(document.document_number, "1")

        versions = list(document.versions.order_by("version_number"))
        self.assertEqual(len(versions), 1)
        version = versions[0]
        self.assertEqual(version.version_number, 1)
        self.assertTrue(version.storage_path_docx)
        self.assertTrue(version.storage_path_docx.endswith(".docx"))
        self.assertTrue(version.storage_path_pdf)
        self.assertTrue(version.storage_path_pdf.endswith(".pdf"))

        # Both files really exist in storage with the right magic bytes.
        from django.core.files.storage import default_storage

        with default_storage.open(version.storage_path_docx, "rb") as fh:
            self.assertTrue(fh.read().startswith(b"PK"))
        with default_storage.open(version.storage_path_pdf, "rb") as fh:
            self.assertEqual(fh.read(), FAKE_PDF)

        items = list(version.line_items.order_by("line_number"))
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].description, "Consulting hours")

        # Exactly one creation audit event for the document...
        self.assertEqual(
            AuditLog.objects.filter(
                action="create",
                entity_type="generated_document",
                entity_id=document.id,
            ).count(),
            1,
        )
        # ...plus the generation events, mirroring the step-wise flow.
        self.assertEqual(
            AuditLog.objects.filter(
                action="generate",
                entity_type="generated_document",
                entity_id=document.id,
            ).count(),
            1,
        )
        self.assertEqual(
            AuditLog.objects.filter(
                action="generate_pdf",
                entity_type="generated_document_version",
                entity_id=version.id,
            ).count(),
            1,
        )

        self.assertEqual(self.counter_value(self.company_a), 1)

    # ------------------------------------------------------------ rollback

    def test_pdf_failure_rolls_back_everything(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        from apps.generated_documents.services.pdf_generator import (
            PdfConversionUnavailableError,
        )

        with mock.patch(
            CONVERTER_PATH, side_effect=PdfConversionUnavailableError("no runtime")
        ):
            response = self.client.post(self.url(), self.payload(), format="json")

        self.assertEqual(
            response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE
        )
        self.assertFalse(GeneratedDocument.objects.exists())
        self.assertFalse(GeneratedDocumentVersion.objects.exists())
        self.assertFalse(DocumentLineItem.objects.exists())
        self.assertFalse(AuditLog.objects.exists())
        self.assertIsNone(self.counter_value(self.company_a))
        self.assertFalse(self.generated_storage_exists())

    def test_docx_failure_rolls_back_everything(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(
            "apps.generated_documents.services.generate_from_input._render_docx",
            side_effect=PdfGenerationError("boom"),
        ):
            response = self.client.post(self.url(), self.payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertNotIn("boom", response.json()["detail"])
        self.assertFalse(GeneratedDocument.objects.exists())
        self.assertFalse(GeneratedDocumentVersion.objects.exists())
        self.assertFalse(DocumentLineItem.objects.exists())
        self.assertFalse(AuditLog.objects.exists())
        self.assertIsNone(self.counter_value(self.company_a))
        self.assertFalse(self.generated_storage_exists())

    def test_failed_attempt_consumes_no_number(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        from apps.generated_documents.services.pdf_generator import (
            PdfConversionUnavailableError,
        )

        with mock.patch(
            CONVERTER_PATH, side_effect=PdfConversionUnavailableError("no runtime")
        ):
            failed = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(failed.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)

        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            first = self.client.post(self.url(), self.payload(), format="json")
            second = self.client.post(self.url(), self.payload(), format="json")

        self.assertEqual(first.status_code, status.HTTP_201_CREATED)
        self.assertEqual(second.status_code, status.HTTP_201_CREATED)
        self.assertEqual(first.json()["document_number"], "1")
        self.assertEqual(second.json()["document_number"], "2")
        self.assertEqual(self.counter_value(self.company_a), 2)

    # ------------------------------------------------------------ security

    def test_unauthenticated_is_rejected(self):
        response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertFalse(GeneratedDocument.objects.exists())

    def test_generation_requires_permission(self):
        self.client.force_authenticate(user=self.alice)
        response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(GeneratedDocument.objects.exists())

    def test_cross_tenant_client_resolves_to_404_and_creates_nothing(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(), self.payload(client_id=self.client_b.id), format="json"
            )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertFalse(GeneratedDocument.objects.exists())
        self.assertIsNone(self.counter_value(self.company_a))

    def test_company_comes_from_authenticated_user(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        document = GeneratedDocument.objects.get(pk=response.json()["id"])
        self.assertEqual(document.company_id, self.company_a.id)

    # ------------------------------------------------------------ validation

    def test_unknown_currency_returns_400_and_creates_nothing(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=FAKE_PDF):
            response = self.client.post(
                self.url(), self.payload(currency_id=999999), format="json"
            )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(GeneratedDocument.objects.exists())
        self.assertIsNone(self.counter_value(self.company_a))

    def test_missing_fields_return_400_and_create_nothing(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        response = self.client.post(self.url(), {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(GeneratedDocument.objects.exists())

    def test_invalid_pdf_output_rolls_back(self):
        self.grant(self.alice, *GENERATE)
        self.client.force_authenticate(user=self.alice)
        with mock.patch(CONVERTER_PATH, return_value=b"not a pdf"):
            response = self.client.post(self.url(), self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(GeneratedDocument.objects.exists())
        self.assertFalse(GeneratedDocumentVersion.objects.exists())
        self.assertIsNone(self.counter_value(self.company_a))
        self.assertFalse(self.generated_storage_exists())
