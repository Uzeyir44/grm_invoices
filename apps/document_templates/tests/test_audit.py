"""
Step 5.2 integration tests: document-template operations produce audit events
(template create/update, version create, version activate) with correct
company, acting user, and payloads.
"""

import uuid
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from docx import Document

from rest_framework import status
from rest_framework.test import APITestCase

from apps.audit_logs.models import AuditLog
from apps.authentication.models import Permission, Role, UserRole
from apps.authentication.tests.test_authentication import make_company, make_user
from apps.reference_data.models import DocumentType

PERMS = (
    "document_templates.read",
    "document_templates.create",
    "document_templates.update",
    "document_templates.version_create",
    "document_templates.version_activate",
)


def build_template_docx():
    document = Document()
    document.add_paragraph("Invoice: {{ document.number }}")
    document.add_paragraph("Client: {{ client.name }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


def build_invalid_template_docx():
    document = Document()
    document.add_paragraph("Bad placeholder: {{ bogus.field }}")
    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out.getvalue()


class TemplateAuditTestCase(APITestCase):
    def setUp(self):
        self.company = make_company()
        self.user = make_user(self.company, email="alice@example.com")
        self.document_type, _ = DocumentType.objects.get_or_create(
            code="invoice", defaults={"name": "Invoice"}
        )
        self.grant(self.user, *PERMS)
        self.auth(self.user)

    def grant(self, user, *codes):
        role = Role.objects.create(name=f"R{uuid.uuid4().hex[:8]}", company=None)
        for code in codes:
            role.permissions.add(Permission.objects.get(code=code))
        UserRole.objects.create(user=user, role=role)

    def auth(self, user):
        self.client.force_authenticate(user=user)

    def create_template(self, name="Invoice Template", document_type=None):
        response = self.client.post(
            reverse("document-template-list-create"),
            {
                "name": name,
                "document_type": document_type or self.document_type.id,
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        return response.json()["id"]

    def upload_file(self, template_id):
        upload = SimpleUploadedFile(
            "tpl.docx", build_template_docx(), content_type="application/octet-stream"
        )
        response = self.client.post(
            reverse("document-template-file", kwargs={"pk": template_id}),
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)


class TemplateCreateAuditTests(TemplateAuditTestCase):
    def test_template_create_records_event(self):
        template_id = self.create_template()
        event = AuditLog.objects.filter(
            entity_type="document_template", entity_id=template_id
        ).get()
        self.assertEqual(event.action, "create")
        self.assertEqual(event.company_id, self.company.id)
        self.assertEqual(event.user_id, self.user.id)
        self.assertEqual(event.new_values["name"], "Invoice Template")
        self.assertEqual(event.new_values["is_active"], True)
        self.assertIsNone(event.previous_values)

    def test_failed_template_create_records_no_event(self):
        before = AuditLog.objects.count()
        response = self.client.post(
            reverse("document-template-list-create"),
            {"name": "X", "document_type": 999999},
            format="json",
        )
        self.assertIn(response.status_code, (status.HTTP_400_BAD_REQUEST, status.HTTP_404_NOT_FOUND))
        self.assertEqual(AuditLog.objects.count(), before)


class TemplateUpdateAuditTests(TemplateAuditTestCase):
    def test_template_update_records_diff(self):
        template_id = self.create_template(name="Before")
        response = self.client.patch(
            reverse("document-template-detail", kwargs={"pk": template_id}),
            {"name": "After"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        events = AuditLog.objects.filter(
            entity_type="document_template", entity_id=template_id
        ).order_by("created_at")
        self.assertEqual(events.count(), 2)
        event = events.last()
        self.assertEqual(event.action, "update")
        self.assertEqual(event.previous_values, {"name": "Before"})
        self.assertEqual(event.new_values, {"name": "After"})
        self.assertEqual(event.company_id, self.company.id)
        self.assertEqual(event.user_id, self.user.id)

    def test_noop_template_update_records_no_event(self):
        template_id = self.create_template(name="Same")
        before = AuditLog.objects.count()
        response = self.client.patch(
            reverse("document-template-detail", kwargs={"pk": template_id}),
            {"name": "Same"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(AuditLog.objects.count(), before)

    def test_failed_template_update_records_no_event(self):
        template_id = self.create_template(name="Before")
        before = AuditLog.objects.count()
        response = self.client.patch(
            reverse("document-template-detail", kwargs={"pk": template_id}),
            {"document_type": 999999},
            format="json",
        )
        self.assertIn(response.status_code, (status.HTTP_400_BAD_REQUEST, status.HTTP_404_NOT_FOUND))
        self.assertEqual(AuditLog.objects.count(), before)


class TemplateVersionAuditTests(TemplateAuditTestCase):
    def setUp(self):
        super().setUp()
        self.template_id = self.create_template()
        self.upload_file(self.template_id)

    def version_list_url(self):
        return reverse(
            "document-template-version-list-create", kwargs={"pk": self.template_id}
        )

    def test_version_create_records_event(self):
        response = self.client.post(self.version_list_url(), {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        version_id = response.json()["id"]
        event = AuditLog.objects.filter(
            entity_type="document_template_version", entity_id=version_id
        ).get()
        self.assertEqual(event.action, "version_create")
        self.assertEqual(event.company_id, self.company.id)
        self.assertEqual(event.user_id, self.user.id)
        self.assertEqual(event.new_values["template_id"], self.template_id)
        self.assertEqual(event.new_values["version_number"], 1)
        self.assertEqual(event.new_values["is_active"], True)

    def test_failed_version_create_records_no_event(self):
        # A template whose file contains an invalid placeholder is rejected at
        # version-creation time, so no version and no audit event are created.
        bad_template_id = self.create_template(
            name="Bad Template",
            document_type=DocumentType.objects.get_or_create(
                code="consulting_contract", defaults={"name": "Consulting Contract"}
            )[0].id,
        )
        upload = SimpleUploadedFile(
            "bad.docx",
            build_invalid_template_docx(),
            content_type="application/octet-stream",
        )
        response = self.client.post(
            reverse("document-template-file", kwargs={"pk": bad_template_id}),
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

        before = AuditLog.objects.count()
        response = self.client.post(
            reverse(
                "document-template-version-list-create",
                kwargs={"pk": bad_template_id},
            ),
            {},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(AuditLog.objects.count(), before)
        self.assertFalse(
            AuditLog.objects.filter(
                entity_type="document_template_version"
            ).exists()
        )

    def test_version_activate_records_event(self):
        first = self.client.post(self.version_list_url(), {}, format="json")
        first_id = first.json()["id"]
        self.client.post(self.version_list_url(), {}, format="json")
        # Version 1 is now inactive (creating version 2 activated it).
        activate_url = reverse(
            "document-template-version-activate",
            kwargs={"pk": self.template_id, "version_id": first_id},
        )
        response = self.client.post(activate_url, {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        event = (
            AuditLog.objects.filter(
                entity_type="document_template_version",
                entity_id=first_id,
                action="version_activate",
            )
            .order_by("created_at")
            .last()
        )
        self.assertIsNotNone(event)
        self.assertEqual(event.previous_values, {"is_active": False})
        self.assertEqual(event.new_values, {"is_active": True})
        self.assertEqual(event.company_id, self.company.id)
        self.assertEqual(event.user_id, self.user.id)